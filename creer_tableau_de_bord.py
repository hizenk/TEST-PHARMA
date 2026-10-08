"""Ajoute l'onglet « Tableau de bord » dans Donnees_propres.xlsx.

Le tableau de bord fonctionne entièrement dans Excel, sans macro : on choisit
les médicaments (Oui / Non), une période et un regroupement, et les chiffres,
le tableau et les graphiques se recalculent tout seuls par formules à partir
des feuilles Pharma_Ventes_Daily et Pharma_Ventes_Hourly.

À relancer après chaque mise à jour des données (le script remplace l'ancien
tableau de bord et ne touche pas aux feuilles de données) :

    python creer_tableau_de_bord.py
    python creer_tableau_de_bord.py --excel autre_classeur.xlsx --sortie copie.xlsx
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ICI = Path(__file__).resolve().parent
EXCEL_PAR_DEFAUT = ICI / "Donnees_propres.xlsx"

FEUILLE = "Tableau de bord"
CALCULS = "Calculs"
JOUR = "Pharma_Ventes_Daily"  # nom de la feuille ET du tableau Excel des ventes journalières
HEURE = "Pharma_Ventes_Hourly"
CODE_ATC = re.compile(r"[A-Z]\d{2}[A-Z]{0,2}")
MAX_LIGNES = 400  # lignes du tableau de résultats

NOMS = {
    "M01AB": "Diclofénac (anti-inflammatoire)", "M01AE": "Ibuprofène (anti-inflammatoire)",
    "N02BA": "Aspirine", "N02BE": "Paracétamol", "N05B": "Anxiolytiques", "N05C": "Hypnotiques, sédatifs",
    "R03": "Asthme, BPCO", "R06": "Antihistaminiques",
}
# Une couleur fixe par médicament (palette validée pour le daltonisme), dans le même ordre partout.
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
REGROUPEMENTS = ["Jour", "Semaine", "Mois", "Année", "Jour de la semaine", "Heure"]
VALEURS = ["Somme", "Moyenne par jour"]
JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août",
        "Septembre", "Octobre", "Novembre", "Décembre"]

# Mise en page
BLEU_FONCE, BLEU, GRIS_TEXTE, GRIS_CLAIR = "1F3864", "2A78D6", "595959", "A6A6A6"
FOND_SAISIE, FOND_TUILE, FOND_ENTETE = "FFF2CC", "F3F6FA", "1F4E78"
FMT_QTE, FMT_DATE, FMT_PCT = "#,##0.00", "dd/mm/yyyy", "0%"
TRAIT = Side(style="thin", color="BFBFBF")
CADRE = Border(left=TRAIT, right=TRAIT, top=TRAIT, bottom=TRAIT)

# Plan de la feuille Tableau de bord
L_PANNEAU = 7          # première ligne des médicaments
L_DEBUT, L_FIN, L_DISPO = None, None, None  # calculées dans construire()
L_ENTETE_RESULTATS = 30
COL_PERIODE = 2        # colonne B


def lettre(col):
    return get_column_letter(col)


def chaine_date(ref):
    """Date au format jj/mm/aaaa, construite sans TEXTE() pour ne pas dépendre de la langue d'Excel."""
    return f'RIGHT("0"&DAY({ref}),2)&"/"&RIGHT("0"&MONTH({ref}),2)&"/"&YEAR({ref})'


def choisir(index, valeurs):
    return f"CHOOSE({index}," + ",".join(f'"{v}"' for v in valeurs) + ")"


def entete(ws, cellule, texte, largeur=1):
    ws[cellule] = texte
    c = ws[cellule]
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=FOND_ENTETE)
    c.alignment = Alignment(vertical="center", indent=1)
    if largeur > 1:
        ligne, col = c.row, c.column
        for j in range(col + 1, col + largeur):
            ws.cell(ligne, j).fill = PatternFill("solid", fgColor=FOND_ENTETE)


def saisie(cellule, valeur=None):
    """Cellule que l'utilisateur modifie : fond jaune, cadre, non verrouillée."""
    if valeur is not None:
        cellule.value = valeur
    cellule.fill = PatternFill("solid", fgColor=FOND_SAISIE)
    cellule.border = CADRE
    cellule.protection = Protection(locked=False)
    cellule.font = Font(bold=True, color=BLEU_FONCE)
    cellule.alignment = Alignment(horizontal="center", vertical="center")


def nommer(wb, nom, ref):
    if nom in wb.defined_names:
        del wb.defined_names[nom]
    wb.defined_names[nom] = DefinedName(nom, attr_text=ref)


def titre_graphique(texte):
    """Titre de graphique sobre (11 pt) plutôt que le titre 14 pt par défaut."""
    police = CharacterProperties(sz=1100, b=True, solidFill=BLEU_FONCE)
    paragraphe = Paragraph(pPr=ParagraphProperties(defRPr=police), r=[RegularTextRun(rPr=police, t=texte)])
    return Title(tx=Text(rich=RichText(p=[paragraphe])), overlay=False)


def axes_visibles(graphe):
    # openpyxl masque les axes par défaut dans les versions récentes d'Excel
    graphe.x_axis.delete = False
    graphe.y_axis.delete = False


def construire(chemin, sortie):
    wb = load_workbook(chemin)
    for nom in (FEUILLE, CALCULS):
        if nom in wb.sheetnames:
            del wb[nom]
    if JOUR not in wb.sheetnames or HEURE not in wb.sheetnames:
        raise SystemExit(f"Feuilles {JOUR} et {HEURE} introuvables dans {chemin}")

    entetes = [c.value for c in wb[JOUR][1]]
    codes = [c for c in entetes if isinstance(c, str) and CODE_ATC.fullmatch(c)]
    for colonne in ["datum", *codes]:
        if colonne not in entetes:
            raise SystemExit(f"Colonne {colonne} introuvable dans {JOUR}")
    for colonne in ["datum", "Hour", *codes]:
        if colonne not in [c.value for c in wb[HEURE][1]]:
            raise SystemExit(f"Colonne {colonne} introuvable dans {HEURE}")
    if len(codes) > len(PALETTE):
        raise SystemExit(f"{len(codes)} médicaments : le tableau de bord en gère {len(PALETTE)} au maximum")

    ws = wb.create_sheet(FEUILLE, 0)
    calc = wb.create_sheet(CALCULS)
    wb.active = 0
    for feuille in wb.worksheets:
        feuille.sheet_view.tabSelected = feuille.title == FEUILLE

    J = lambda col: f"{JOUR}[{col}]"
    H = lambda col: f"{HEURE}[{col}]"
    dans_periode = f'{J("datum")},">="&Debut,{J("datum")},"<="&Fin'

    # ------------------------------------------------------------------ Titre et mode d'emploi
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws["B2"] = "Ventes de la pharmacie : tableau de bord"
    ws["B2"].font = Font(bold=True, size=20, color=BLEU_FONCE)
    ws["B3"] = ("Choisissez les médicaments, la période et le regroupement dans les cellules jaunes : "
                "tout se met à jour automatiquement.")
    ws["B3"].font = Font(italic=True, color=GRIS_TEXTE)

    # ------------------------------------------------------------------ ① Médicaments
    entete(ws, "B5", "① Médicaments", 5)
    for j, titre in enumerate(["Médicament", "Inclure", "Total période", "Moyenne / jour", "Part"], start=2):
        c = ws.cell(6, j, titre)
        c.font = Font(bold=True, color=GRIS_TEXTE, size=9)
        c.alignment = Alignment(horizontal="left" if j == 2 else "center")
        c.border = Border(bottom=TRAIT)
    lignes_med = {}
    for i, code in enumerate(codes):
        r = L_PANNEAU + i
        lignes_med[code] = r
        ws.cell(r, 2, f"{code} · {NOMS.get(code, code)}").font = Font(color="000000")
        ws.cell(r, 2).border = Border(left=Side(style="thick", color=PALETTE[i]))
        saisie(ws.cell(r, 3), "Oui")
        ws.cell(r, 4, f"=SUMIFS({J(code)},{dans_periode})").number_format = FMT_QTE
        ws.cell(r, 5, f"=IF(NbJoursPeriode>0,D{r}/NbJoursPeriode,0)").number_format = FMT_QTE
        ws.cell(r, 6, f'=IF(AND(C{r}="Oui",TotalSelection>0),D{r}/TotalSelection,"")').number_format = FMT_PCT
    l_total = L_PANNEAU + len(codes)
    plage_inclure = f"$C${L_PANNEAU}:$C${l_total - 1}"
    ws.cell(l_total, 2, "Total de la sélection").font = Font(bold=True)
    ws.cell(l_total, 4, f'=SUMPRODUCT(({plage_inclure}="Oui")*$D${L_PANNEAU}:$D${l_total - 1})')
    ws.cell(l_total, 5, f"=IF(NbJoursPeriode>0,D{l_total}/NbJoursPeriode,0)")
    for j in range(2, 7):
        c = ws.cell(l_total, j)
        c.font = Font(bold=True)
        c.border = Border(top=TRAIT)
        if j in (4, 5):
            c.number_format = FMT_QTE
    nommer(wb, "TotalSelection", f"'{FEUILLE}'!$D${l_total}")

    # ------------------------------------------------------------------ Mode d'emploi
    entete(ws, "H5", "Mode d'emploi", 5)
    consignes = [
        "1. Mettez Oui ou Non devant chaque médicament.",
        "2. Saisissez les dates de début et de fin.",
        "3. Choisissez le regroupement et la valeur.",
        "4. Lisez les indicateurs, le tableau et les graphiques.",
        "",
        "Seules les cellules jaunes sont modifiables.",
        "Données lues dans Pharma_Ventes_Daily et",
        "Pharma_Ventes_Hourly : rien à recopier.",
    ]
    for k, texte in enumerate(consignes):
        ws.cell(6 + k, 8, texte).font = Font(size=10, color=GRIS_TEXTE, bold=k < 4)

    # ------------------------------------------------------------------ ② Période
    l = l_total + 2
    entete(ws, f"B{l}", "② Période", 5)
    l_debut, l_fin, l_dispo = l + 1, l + 2, l + 3
    ws[f"B{l_debut}"], ws[f"B{l_fin}"] = "Du", "Au"
    saisie(ws[f"C{l_debut}"], datetime(2018, 1, 1))
    saisie(ws[f"C{l_fin}"], datetime(2018, 12, 31))
    for cellule in (f"C{l_debut}", f"C{l_fin}"):
        ws[cellule].number_format = FMT_DATE
        ws.merge_cells(f"{cellule}:{cellule.replace('C', 'D')}")
    ws[f"B{l_dispo}"] = "Données disponibles"
    ws[f"C{l_dispo}"] = f"=MIN({J('datum')})"
    ws[f"D{l_dispo}"] = f"=MAX({J('datum')})"
    for cellule in (f"B{l_dispo}", f"C{l_dispo}", f"D{l_dispo}"):
        ws[cellule].font = Font(size=9, color=GRIS_TEXTE)
        ws[cellule].number_format = FMT_DATE
    nommer(wb, "Debut", f"'{FEUILLE}'!$C${l_debut}")
    nommer(wb, "Fin", f"'{FEUILLE}'!$C${l_fin}")

    # ------------------------------------------------------------------ ③ Affichage
    l = l_dispo + 2
    entete(ws, f"B{l}", "③ Affichage", 5)
    l_regroup, l_valeur = l + 1, l + 2
    ws[f"B{l_regroup}"], ws[f"B{l_valeur}"] = "Regrouper par", "Valeur"
    saisie(ws[f"C{l_regroup}"], "Mois")
    saisie(ws[f"C{l_valeur}"], "Somme")
    for cellule in (f"C{l_regroup}", f"C{l_valeur}"):
        ws.merge_cells(f"{cellule}:{cellule.replace('C', 'D')}")
    nommer(wb, "Regroupement", f"'{FEUILLE}'!$C${l_regroup}")
    nommer(wb, "Valeur", f"'{FEUILLE}'!$C${l_valeur}")
    l_message = l_valeur + 1
    ws[f"B{l_message}"] = (
        f'=IF(Fin<Debut,"⚠ La date de fin est avant la date de début.",'
        f'IF(COUNTIF({plage_inclure},"Oui")=0,"⚠ Aucun médicament n\'est inclus.",'
        f'IF(NbLignes>{MAX_LIGNES},"⚠ "&NbLignes&" lignes : seules les {MAX_LIGNES} premières sont affichées, '
        f'choisissez un regroupement plus large.","")))'
    )
    ws[f"B{l_message}"].font = Font(bold=True, color="C00000")
    assert l_message < L_ENTETE_RESULTATS - 1, "le panneau de saisie déborde sur le tableau de résultats"

    # Listes déroulantes et contrôle des dates
    dv_oui = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False,
                            error="Choisissez Oui ou Non.", errorTitle="Valeur non valide")
    dv_oui.add(plage_inclure.replace("$", ""))
    dv_dates = DataValidation(type="date", operator="between", formula1=f"$C${l_dispo}", formula2=f"$D${l_dispo}",
                              error="Choisissez une date comprise dans les données disponibles.",
                              errorTitle="Date hors des données", showErrorMessage=True)
    dv_dates.add(f"C{l_debut}")
    dv_dates.add(f"C{l_fin}")
    dv_regroup = DataValidation(type="list", formula1='"' + ",".join(REGROUPEMENTS) + '"', allow_blank=False,
                                error="Choisissez un regroupement dans la liste.", errorTitle="Valeur non valide")
    dv_regroup.add(f"C{l_regroup}")
    dv_valeur = DataValidation(type="list", formula1='"' + ",".join(VALEURS) + '"', allow_blank=False,
                               error="Choisissez une valeur dans la liste.", errorTitle="Valeur non valide")
    dv_valeur.add(f"C{l_valeur}")
    for dv in (dv_oui, dv_dates, dv_regroup, dv_valeur):
        dv.showErrorMessage = True
        ws.add_data_validation(dv)

    # Médicaments exclus en gris
    ws.conditional_formatting.add(
        f"B{L_PANNEAU}:F{l_total - 1}",
        FormulaRule(formula=[f'$C{L_PANNEAU}="Non"'], font=Font(color=GRIS_CLAIR), stopIfTrue=False))
    ws.conditional_formatting.add(
        f"F{L_PANNEAU}:F{l_total - 1}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=BLEU, showValue=True))

    # ------------------------------------------------------------------ Feuille de calculs (masquée)
    # Ventes journalières de la sélection (somme des médicaments marqués Oui) et filtre de période
    selection = "(" + "+".join(f'{J(code)}*(\'{FEUILLE}\'!$C${lignes_med[code]}="Oui")' for code in codes) + ")"
    periode = f'({J("datum")}>=Debut)*({J("datum")}<=Fin)'
    calc["A1"], calc["B1"] = "Paramètre", "Valeur"
    parametres = [
        ("N° de regroupement", "=IFERROR(MATCH(Regroupement,{" + ",".join(f'"{r}"' for r in REGROUPEMENTS) + "},0),3)",
         "NumRegroupement"),
        ("Lundi de la 1re semaine", "=Debut-WEEKDAY(Debut,2)+1", "LundiDebut"),
        ("Nombre de lignes", "=IF(Fin<Debut,0,CHOOSE(NumRegroupement,Fin-Debut+1,INT((Fin-LundiDebut)/7)+1,"
                             "(YEAR(Fin)-YEAR(Debut))*12+MONTH(Fin)-MONTH(Debut)+1,YEAR(Fin)-YEAR(Debut)+1,7,24))",
         "NbLignes"),
        ("Lignes affichées", f"=MIN(NbLignes,{MAX_LIGNES})", "NbAffichees"),
        ("Moyenne par jour ?", '=Valeur="Moyenne par jour"', "ParJour"),
        ("Jours de données", f"=COUNTIFS({dans_periode})", "NbJoursPeriode"),
        ("Meilleur jour (quantité)", f"=SUMPRODUCT(MAX({selection}*{periode}))", "MaxJour"),
    ]
    for i, (libelle, formule, nom) in enumerate(parametres, start=2):
        calc.cell(i, 1, libelle)
        calc.cell(i, 2, formule)
        nommer(wb, nom, f"{CALCULS}!$B${i}")

    L0 = 10  # première ligne des périodes dans Calculs
    for j, titre in enumerate(["N°", "Affichée", "Début", "Fin", "Début (borné)", "Fin (bornée)", "Libellé",
                               "X graphique", "Y graphique"], start=1):
        calc.cell(L0 - 1, j, titre).font = Font(bold=True)
    libelle_jour = choisir("WEEKDAY(C{r},2)", ["lun.", "mar.", "mer.", "jeu.", "ven.", "sam.", "dim."])
    for i in range(1, MAX_LIGNES + 1):
        r = L0 + i - 1
        calc.cell(r, 1, i)
        calc.cell(r, 2, f"=A{r}<=NbAffichees")
        calc.cell(r, 3, f'=IF(NOT(B{r}),"",CHOOSE(NumRegroupement,Debut+A{r}-1,LundiDebut+7*(A{r}-1),'
                        f'DATE(YEAR(Debut),MONTH(Debut)+A{r}-1,1),DATE(YEAR(Debut)+A{r}-1,1,1),A{r},A{r}-1))')
        calc.cell(r, 4, f'=IF(NOT(B{r}),"",CHOOSE(NumRegroupement,C{r},C{r}+6,DATE(YEAR(C{r}),MONTH(C{r})+1,0),'
                        f'DATE(YEAR(C{r}),12,31),C{r},C{r}))')
        calc.cell(r, 5, f'=IF(NOT(B{r}),"",IF(NumRegroupement<=4,MAX(C{r},Debut),C{r}))')
        calc.cell(r, 6, f'=IF(NOT(B{r}),"",IF(NumRegroupement<=4,MIN(D{r},Fin),D{r}))')
        calc.cell(r, 7, f'=IF(NOT(B{r}),"",CHOOSE(NumRegroupement,'
                        f'{libelle_jour.format(r=r)}&" "&{chaine_date(f"C{r}")},'
                        f'"Semaine du "&{chaine_date(f"C{r}")},'
                        f'{choisir(f"MONTH(C{r})", MOIS)}&" "&YEAR(C{r}),'
                        f'""&YEAR(C{r}),'
                        f'{choisir(f"A{r}", JOURS)},'
                        f'RIGHT("0"&(A{r}-1),2)&"h"))')
        rd = L_ENTETE_RESULTATS + i
        calc.cell(r, 8, f"=IF(AND(B{r},NumRegroupement<=4),C{r},NA())").number_format = FMT_DATE
        calc.cell(r, 9, f"=IF(AND(B{r},NumRegroupement<=4,ISNUMBER('{FEUILLE}'!K{rd})),'{FEUILLE}'!K{rd},NA())")

    # Profils fixes : jour de la semaine (moyenne par jour) et heure (moyenne par jour), par médicament
    L_PROFILS = L0 + MAX_LIGNES + 2
    calc.cell(L_PROFILS - 1, 1, "Jour").font = Font(bold=True)
    calc.cell(L_PROFILS - 1, 2, "Moyenne / jour").font = Font(bold=True)
    for k, jour in enumerate(JOURS, start=1):
        r = L_PROFILS + k - 1
        meme_jour = f"(WEEKDAY({J('datum')},2)={k})"
        calc.cell(r, 1, jour[:3] + ".")
        calc.cell(r, 2, f"=IFERROR(SUMPRODUCT({meme_jour}*{periode}*{selection})"
                        f"/SUMPRODUCT({meme_jour}*{periode}),0)")
    L_HEURES = L_PROFILS + 9
    calc.cell(L_HEURES - 1, 1, "Heure").font = Font(bold=True)
    calc.cell(L_HEURES - 1, 2, "Moyenne / jour").font = Font(bold=True)
    for h in range(24):
        r = L_HEURES + h
        criteres = f'{H("Hour")},{h},{H("datum")},">="&Debut,{H("datum")},"<"&Fin+1'
        somme = "+".join(f'(\'{FEUILLE}\'!$C${lignes_med[code]}="Oui")*SUMIFS({H(code)},{criteres})' for code in codes)
        calc.cell(r, 1, f"{h:02d}h")
        calc.cell(r, 2, f"=IFERROR(({somme})/COUNTIFS({criteres}),0)")
    L_PARTS = L_HEURES + 26
    calc.cell(L_PARTS - 1, 1, "Médicament").font = Font(bold=True)
    calc.cell(L_PARTS - 1, 2, "Total période").font = Font(bold=True)
    for i, code in enumerate(codes):
        r = L_PARTS + i
        calc.cell(r, 1, code)
        calc.cell(r, 2, f"=IF('{FEUILLE}'!C{lignes_med[code]}=\"Oui\",'{FEUILLE}'!D{lignes_med[code]},0)")
    calc.column_dimensions["A"].width = 24
    calc.column_dimensions["G"].width = 26
    calc.sheet_state = "hidden"

    # ------------------------------------------------------------------ Indicateurs (tuiles)
    tuiles = [
        ("N", "Quantité vendue", "=TotalSelection", FMT_QTE, '="sur "&NbJoursPeriode&" jours"'),
        ("Q", "Moyenne par jour", "=IF(NbJoursPeriode>0,TotalSelection/NbJoursPeriode,0)", FMT_QTE,
         '="pour la sélection"'),
        ("T", "Meilleur jour", None, FMT_DATE, None),
        ("W", "Jours de données", "=NbJoursPeriode", "0", '="du "&' + chaine_date("Debut") + '&" au "&'
         + chaine_date("Fin")),
    ]
    for col, libelle, formule, fmt, sous_titre in tuiles:
        c0 = ws[f"{col}5"].column
        zone = [ws.cell(r, j) for r in range(5, 9) for j in range(c0, c0 + 3)]
        for c in zone:
            c.fill = PatternFill("solid", fgColor=FOND_TUILE)
        ws.cell(5, c0, libelle).font = Font(size=9, bold=True, color=GRIS_TEXTE)
        ws.merge_cells(start_row=6, start_column=c0, end_row=7, end_column=c0 + 2)
        valeur = ws.cell(6, c0)
        valeur.font = Font(size=20, bold=True, color=BLEU_FONCE)
        valeur.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        valeur.number_format = fmt
        if formule:
            valeur.value = formule
        if sous_titre:
            ws.cell(8, c0, sous_titre).font = Font(size=9, color=GRIS_TEXTE)
    # Meilleur jour : jour de la période où la sélection s'est le plus vendue
    ws["T6"] = f'=IF(MaxJour<=0,"—",INDEX({J("datum")},MATCH(MaxJour,INDEX({selection}*{periode},0),0)))'
    ws["T8"] = '=IF(MaxJour<=0,"","avec "&FIXED(MaxJour,2)&" ventes")'
    ws["T8"].font = Font(size=9, color=GRIS_TEXTE)

    # ------------------------------------------------------------------ Tableau des résultats
    l = L_ENTETE_RESULTATS - 1
    ws[f"B{l}"] = '="④ Résultats par "&LOWER(Regroupement)&IF(ParJour," (moyenne par jour)"," (somme)")'
    ws[f"B{l}"].font = Font(bold=True, color="FFFFFF", size=11)
    colonnes = ["Période", *codes, "Total sélection", "Nb jours"]
    col_total, col_nb = 3 + len(codes), 4 + len(codes)
    for j in range(2, col_nb + 1):
        ws.cell(l, j).fill = PatternFill("solid", fgColor=FOND_ENTETE)
    for j, titre in enumerate(colonnes, start=2):
        c = ws.cell(L_ENTETE_RESULTATS, j, titre)
        c.font = Font(bold=True, color=GRIS_TEXTE, size=9)
        c.alignment = Alignment(horizontal="left" if j == 2 else "right")
        c.border = Border(bottom=TRAIT)
    lt, ln = lettre(col_total), lettre(col_nb)
    for i in range(1, MAX_LIGNES + 1):
        rc, rd = L0 + i - 1, L_ENTETE_RESULTATS + i
        actif = f"{CALCULS}!$B{rc}"
        debut_l, fin_l, num, cle = (f"{CALCULS}!$E{rc}", f"{CALCULS}!$F{rc}", f"{CALCULS}!$A{rc}", f"{CALCULS}!$C{rc}")
        ws.cell(rd, 2, f'=IF({actif},{CALCULS}!$G{rc},"")')
        jours_ligne = (f'COUNTIFS({J("datum")},">="&{debut_l},{J("datum")},"<="&{fin_l})',
                       f'SUMPRODUCT((WEEKDAY({J("datum")},2)={num})*{periode})',
                       f'COUNTIFS({H("Hour")},{cle},{H("datum")},">="&Debut,{H("datum")},"<"&Fin+1)')
        ws.cell(rd, col_nb, f'=IF({actif},IF(NumRegroupement<=4,{jours_ligne[0]},'
                                f'IF(NumRegroupement=5,{jours_ligne[1]},{jours_ligne[2]})),"")').number_format = "0"
        for k, code in enumerate(codes):
            ventes = (f'SUMIFS({J(code)},{J("datum")},">="&{debut_l},{J("datum")},"<="&{fin_l})',
                      f'SUMPRODUCT((WEEKDAY({J("datum")},2)={num})*{periode}*{J(code)})',
                      f'SUMIFS({H(code)},{H("Hour")},{cle},{H("datum")},">="&Debut,{H("datum")},"<"&Fin+1)')
            formule = (f'=IF(AND({actif},$C${lignes_med[code]}="Oui"),'
                       f'IF(NumRegroupement<=4,{ventes[0]},IF(NumRegroupement=5,{ventes[1]},{ventes[2]}))'
                       f'/IF(ParJour,MAX(1,${ln}{rd}),1),"")')
            ws.cell(rd, 3 + k, formule).number_format = FMT_QTE
        ws.cell(rd, col_total, f'=IF({actif},SUM(C{rd}:{lettre(col_total - 1)}{rd}),"")').number_format = FMT_QTE
        ws.cell(rd, col_total).font = Font(bold=True)
    derniere = L_ENTETE_RESULTATS + MAX_LIGNES
    ws.conditional_formatting.add(
        f"{lt}{L_ENTETE_RESULTATS + 1}:{lt}{derniere}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=BLEU, showValue=True))
    for k, code in enumerate(codes):  # en-tête des médicaments exclus en gris
        cellule = f"{lettre(3 + k)}{L_ENTETE_RESULTATS}"
        ws.conditional_formatting.add(cellule, FormulaRule(
            formula=[f'$C${lignes_med[code]}="Non"'], font=Font(color=GRIS_CLAIR)))
    ws.conditional_formatting.add(  # une ligne sur deux légèrement grisée pour la lecture
        f"B{L_ENTETE_RESULTATS + 1}:{ln}{derniere}",
        FormulaRule(formula=[f'AND($B{L_ENTETE_RESULTATS + 1}<>"",MOD(ROW(),2)=0)'],
                    fill=PatternFill("solid", fgColor="F7F7F7")))

    # ------------------------------------------------------------------ Graphiques
    evolution = ScatterChart()
    evolution.title = titre_graphique("Évolution de la sélection (regroupement par jour, semaine, mois ou année)")
    evolution.style = 2
    evolution.legend = None
    evolution.display_blanks = "gap"
    serie = Series(Reference(calc, min_col=9, min_row=L0, max_row=L0 + MAX_LIGNES - 1),
                   Reference(calc, min_col=8, min_row=L0, max_row=L0 + MAX_LIGNES - 1), title="Sélection")
    serie.graphicalProperties.line.solidFill = BLEU
    serie.graphicalProperties.line.width = 25400  # 2 pt
    serie.marker = Marker(symbol="circle", size=5)
    serie.marker.graphicalProperties = GraphicalProperties(solidFill=BLEU)
    serie.marker.graphicalProperties.line.solidFill = BLEU
    serie.smooth = False
    evolution.series.append(serie)
    evolution.x_axis.number_format = "dd/mm/yy"
    evolution.x_axis.majorGridlines = None
    evolution.y_axis.number_format = "#,##0"
    evolution.y_axis.title = "Quantité"
    evolution.y_axis.scaling.min = 0
    axes_visibles(evolution)
    evolution.width, evolution.height = 22.5, 9.5
    ws.add_chart(evolution, "N10")

    def barres(intitule, ligne, n, couleurs=None, horizontal=False):
        g = BarChart()
        g.type = "bar" if horizontal else "col"
        g.title = titre_graphique(intitule)
        g.style = 2
        g.legend = None
        g.gapWidth = 40
        g.add_data(Reference(calc, min_col=2, min_row=ligne - 1, max_row=ligne + n - 1), titles_from_data=True)
        g.set_categories(Reference(calc, min_col=1, min_row=ligne, max_row=ligne + n - 1))
        s = g.series[0]
        s.graphicalProperties.solidFill = BLEU
        s.graphicalProperties.line.noFill = True
        for idx, couleur in enumerate(couleurs or []):
            point = DataPoint(idx=idx)
            point.graphicalProperties.solidFill = couleur
            point.graphicalProperties.line.noFill = True
            s.dPt.append(point)
        g.y_axis.number_format = "#,##0.0"
        g.y_axis.scaling.min = 0  # des barres partent toujours de zéro
        g.y_axis.majorGridlines = None if horizontal else g.y_axis.majorGridlines
        axes_visibles(g)
        return g

    profil_semaine = barres("Ventes moyennes selon le jour de la semaine", L_PROFILS, 7)
    profil_semaine.width, profil_semaine.height = 11, 7.5
    ws.add_chart(profil_semaine, "N29")
    profil_heure = barres("Ventes moyennes par jour selon l'heure", L_HEURES, 24)
    profil_heure.width, profil_heure.height = 11, 7.5
    ws.add_chart(profil_heure, "T29")
    parts = barres("Total de la période par médicament", L_PARTS, len(codes), PALETTE[:len(codes)], horizontal=True)
    parts.x_axis.scaling.orientation = "maxMin"  # même ordre que la liste des médicaments
    parts.series[0].dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                          showLegendKey=False, showPercent=False, showLeaderLines=False)
    parts.series[0].dLbls.numFmt = "#,##0;-#,##0;;"  # rien d'affiché pour les médicaments exclus (0)
    parts.width, parts.height = 22.5, 7.5
    ws.add_chart(parts, "N45")

    # ------------------------------------------------------------------ Largeurs, protection, recalcul
    largeurs = {"A": 2, "B": 30, lt: 14, ln: 9, "M": 3}
    for k in range(len(codes)):
        largeurs[lettre(3 + k)] = 12
    for col in range(14, 26):
        largeurs.setdefault(lettre(col), 9.5)
    for col, largeur in largeurs.items():
        ws.column_dimensions[col].width = largeur
    ws.row_dimensions[2].height = 30

    # Impression : le tableau de bord tient sur une page paysage (le détail des résultats suit)
    ws.print_area = "A1:Y60"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1

    ws.protection.sheet = True  # sans mot de passe : Révision > Ôter la protection pour modifier
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(sortie)
    print(f"Tableau de bord ajouté : {sortie} ({len(codes)} médicaments, {MAX_LIGNES} lignes de résultats)")


def main():
    parser = argparse.ArgumentParser(description="Ajoute l'onglet Tableau de bord dans le classeur des ventes.")
    parser.add_argument("--excel", type=Path, default=EXCEL_PAR_DEFAUT, help="classeur à compléter")
    parser.add_argument("--sortie", type=Path, help="classeur à écrire (défaut : remplace --excel)")
    args = parser.parse_args()
    construire(args.excel, args.sortie or args.excel)


if __name__ == "__main__":
    main()
